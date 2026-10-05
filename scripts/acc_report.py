#!/usr/bin/env python3
"""Hyperscout Holding books: report builder.

Builds, from the ledger tabs read with get_values (displayed values):
  --type month    --period 2026-09            -> Maandoverzicht PDF (for Jan, English)
  --type quarter  --period 2026-Q3 [--draft]  -> Kwartaaloverzicht PDF (for Jan) and
                                                the Excel file for Serdal (Dutch)

data.json:
  {"boekingen":   rows of Boekingen!A1:L       (header row included),
   "facturen_in": rows of 'Facturen in'!A1:T  (header row included),
   "facturen_uit":rows of 'Facturen uit'!A1:R (header row included),
   "amex_owed":   'Amex declaraties'!B5 (displayed value),
   "notes":       ["optional line for Jan", ...]}

Usage:
  python3 acc_report.py --type month --period 2026-09 --data data.json --out /path
  python3 acc_report.py --type quarter --period 2026-Q3 --data data.json --out /path [--draft]
Prints the paths of the files it wrote.
"""
import argparse, json, os, re
from datetime import date

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether

DB = colors.HexColor("#14213D"); BEIGE = colors.HexColor("#EFE8DC"); LINE = colors.HexColor("#C9C2B5")
GREY = colors.HexColor("#5F6368"); RED = colors.HexColor("#DE2E26")
LOGO = [(911,1),(795,1),(774,4),(739,22),(712,57),(706,78),(706,477),(701,490),(683,502),(619,502),(597,487),(450,61),(437,37),(417,18),(391,5),(365,1),(89,1),(50,11),(20,34),(6,59),(0,90),(0,904),(6,935),(34,972),(61,988),(80,992),(516,992),(549,980),(577,956),(594,917),(594,660),(606,645),(615,645),(623,652),(845,979),(875,996),(913,1000),(954,989),(981,967),(999,926),(999,78),(987,45),(966,22),(945,9)]
VAT_HIGH = 0.21
FULL = ["January", "February", "March", "April", "May", "June", "July", "August",
        "September", "October", "November", "December"]

# ------------------------------------------------------------------ parsing
def num(s):
    if s is None: return 0.0
    if isinstance(s, (int, float)): return float(s)
    s = str(s).strip()
    if s in ("", "-"): return 0.0
    neg = s.startswith("(") or s.startswith("-") or s.startswith("−")
    s = re.sub(r"[^0-9.,]", "", s).replace(",", "")
    if not s: return 0.0
    return -float(s) if neg else float(s)

def rows(data, key, width):
    out = []
    for r in (data.get(key) or [])[1:]:
        r = (list(r) + [""] * width)[:width]
        if r[0]: out.append(r)
    return out

def quarter_of(d):          # "2026-09-15" -> "2026-Q3"
    return f"{d[:4]}-Q{(int(d[5:7]) - 1) // 3 + 1}"

def q_months(q):            # "2026-Q3" -> ["2026-07","2026-08","2026-09"]
    y, n = q.split("-Q"); n = int(n)
    return [f"{y}-{m:02d}" for m in range((n - 1) * 3 + 1, n * 3 + 1)]

def deadline(q):
    y, n = map(int, q.split("-Q")); m = n * 3 + 1
    if m > 12: y, m = y + 1, 1
    nxt = date(y + (m == 12), (m % 12) + 1, 1)
    return date.fromordinal(nxt.toordinal() - 1)

def eur(v, dash=True):
    if abs(v) < 0.005: return "-" if dash else "€ 0.00"
    s = f"€ {abs(v):,.2f}"
    return f"({s})" if v < 0 else s

def esc(t): return str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

# ------------------------------------------------------------------ VAT boxes (Python check, same rules as the ledger)
def vat_boxes(fin, fuit, q):
    b = {k: 0.0 for k in ("1a", "1a_btw", "1b", "1b_btw", "3a", "3b", "4a", "4a_btw", "4b", "4b_btw", "5b", "buiten EU")}
    for r in fuit:
        if quarter_of(r[0]) != q: continue
        box = r[9]
        if box in ("1a", "1b"): b[box] += num(r[5]); b[box + "_btw"] += num(r[7])
        elif box in b: b[box] += num(r[5])
    for r in fin:
        if quarter_of(r[0]) != q: continue
        rc = round(num(r[5]) * VAT_HIGH, 2) if str(r[6]).startswith("verlegd") else 0.0
        if r[9] in ("4a", "4b"): b[r[9]] += num(r[5]); b[r[9] + "_btw"] += rc
        if r[10] == "Ja": b["5b"] += num(r[7]) + rc
    b["5a"] = b["1a_btw"] + b["1b_btw"] + b["4a_btw"] + b["4b_btw"]
    b["total"] = b["5a"] - b["5b"]
    return b

# ------------------------------------------------------------------ PDF
def S():
    return dict(h1=ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=15, textColor=DB, leading=19, spaceAfter=2 * mm),
                h2=ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=11, textColor=DB, spaceBefore=5 * mm, spaceAfter=2 * mm),
                body=ParagraphStyle("b", fontName="Helvetica", fontSize=9, leading=12.5),
                small=ParagraphStyle("s", fontName="Helvetica", fontSize=7.5, leading=10, textColor=GREY))

def table(header, body, widths, money_from=1, bold_last=False):
    data = [header] + body
    st = [("FONT", (0, 0), (-1, -1), "Helvetica", 8.6), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.4),
          ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("BACKGROUND", (0, 0), (-1, 0), DB),
          ("ALIGN", (money_from, 0), (-1, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
          ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
          ("LINEBELOW", (0, 1), (-1, -1), 0.3, LINE)]
    if bold_last and body:
        st += [("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 8.6), ("BACKGROUND", (0, -1), (-1, -1), BEIGE)]
    t = Table(data, colWidths=widths, repeatRows=1); t.setStyle(TableStyle(st)); return t

def kpis(items):
    cells = [[Paragraph(f'<font size="7.5" color="#5F6368">{l.upper()}</font><br/><font size="13" color="#14213D"><b>{v}</b></font>',
                        ParagraphStyle("k", leading=16)) for l, v in items]]
    t = Table(cells, colWidths=[174 * mm / len(items)] * len(items))
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), BEIGE), ("TOPPADDING", (0, 0), (-1, -1), 6),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 7), ("LEFTPADDING", (0, 0), (-1, -1), 8),
                           ("LINEAFTER", (0, 0), (-2, -1), 0.6, colors.white)]))
    return t

def build_pdf(path, title, subtitle, story_body, today):
    W, H = A4
    def deco(c, doc):
        c.saveState()
        c.setFillColor(DB); c.rect(0, H - 24 * mm, W, 24 * mm, stroke=0, fill=1)
        c.setFillColor(colors.white); p = c.beginPath(); sx, sy = 13.4 * mm / 1000, 14 * mm / 1000
        p.moveTo(18 * mm + LOGO[0][0] * sx, H - 19 * mm + LOGO[0][1] * sy)
        for x, y in LOGO[1:]: p.lineTo(18 * mm + x * sx, H - 19 * mm + y * sy)
        p.close(); c.drawPath(p, fill=1, stroke=0)
        c.setFont("Helvetica-Bold", 15); c.drawString(35 * mm, H - 13 * mm, "Hyperscout")
        c.setFont("Helvetica", 8.5); c.drawString(35 * mm, H - 18 * mm, "Holding books · bookkeeping")
        c.setFont("Helvetica-Bold", 9); c.drawRightString(W - 18 * mm, H - 12.5 * mm, title)
        c.setFont("Helvetica", 8.5); c.drawRightString(W - 18 * mm, H - 17.5 * mm, f"Issued {today.strftime('%d %B %Y')}")
        c.setStrokeColor(LINE); c.setLineWidth(0.4); c.line(18 * mm, 14 * mm, W - 18 * mm, 14 * mm)
        c.setFillColor(GREY); c.setFont("Helvetica", 7)
        c.drawString(18 * mm, 9.5 * mm, "Hyperscout Holding B.V. · confidential, for Jan only · not the statutory accounts")
        c.drawRightString(W - 18 * mm, 9.5 * mm, f"Page {doc.page}")
        c.restoreState()
    st = S()
    story = [Paragraph(title, st["h1"]), Paragraph(subtitle, st["body"]), Spacer(1, 4 * mm)] + story_body
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=31 * mm,
                            bottomMargin=20 * mm, title=title, author="Hyperscout Accountant NL")
    doc.build(story, onFirstPage=deco, onLaterPages=deco)

def missing_rows(boek, upto=None):
    out = []
    for r in sorted(boek, key=lambda r: r[0]):
        if r[8] == "Factuur ontbreekt" and (upto is None or r[0][:7] <= upto):
            out.append([r[0], r[1], esc(r[3])[:38], esc(r[4])[:40], eur(num(r[5]))])
    return out

def vat_table(b):
    body = [["1a  Sales, standard rate", eur(b["1a"]), eur(b["1a_btw"])],
            ["1b  Sales, reduced rate", eur(b["1b"]), eur(b["1b_btw"])],
            ["3a  Goods to outside the EU", eur(b["3a"]), ""],
            ["3b  Services and goods to EU businesses (ICP)", eur(b["3b"]), ""],
            ["4a  Services from outside the EU", eur(b["4a"]), eur(b["4a_btw"])],
            ["4b  Services and goods from the EU", eur(b["4b"]), eur(b["4b_btw"])],
            ["5a  BTW owed", "", eur(b["5a"])],
            ["5b  BTW deductible (voorbelasting)", "", eur(b["5b"])],
            ["To pay (+) or to get back (−)", "", eur(b["total"])]]
    return table(["Box", "Base", "BTW"], body, [114 * mm, 30 * mm, 30 * mm], bold_last=True)

def month_pdf(a, data, today):
    fin, fuit, boek = rows(data, "facturen_in", 20), rows(data, "facturen_uit", 18), rows(data, "boekingen", 12)
    m = a.period; q = quarter_of(m + "-01")
    in_month = [r for r in boek if r[0][:7] == m]
    miss = missing_rows(boek, upto=m)
    b = vat_boxes(fin, fuit, q)
    st = S(); body = []
    body.append(kpis([("Missing invoices", str(len(miss))), (f"BTW {q} so far", eur(b["total"], False)),
                      ("Owed to Jan (Amex)", eur(num(data.get("amex_owed")), False)), ("Lines booked", str(len(in_month)))]))
    for n in data.get("notes") or []:
        body += [Spacer(1, 2 * mm), Paragraph(esc(n), st["body"])]
    if miss:
        body.append(KeepTogether([Paragraph("Missing invoices: please find these", st["h2"]),
                                  table(["Date", "Source", "Counterparty", "Description", "Amount"], miss,
                                        [22 * mm, 16 * mm, 50 * mm, 60 * mm, 26 * mm], money_from=4)]))
    else:
        body += [Paragraph("Missing invoices", st["h2"]), Paragraph("None. Every booking has its invoice.", st["body"])]
    body.append(KeepTogether([Paragraph(f"BTW position {q} so far (due {deadline(q).strftime('%d %B %Y')})", st["h2"]), vat_table(b)]))
    per = {}
    for r in in_month: per[r[6] or "Not classified"] = per.get(r[6] or "Not classified", 0.0) + num(r[5])
    if per:
        tb = [[esc(k), eur(v)] for k, v in sorted(per.items(), key=lambda x: x[1])]
        tb.append(["Net movement", eur(sum(per.values()))])
        body.append(KeepTogether([Paragraph(f"Bookings in {FULL[int(m[5:]) - 1]} by ledger account (gross, as on bank and card)", st["h2"]),
                                  table(["Ledger account", "Amount"], tb, [134 * mm, 40 * mm], bold_last=True)]))
    body += [Spacer(1, 3 * mm), Paragraph("BTW is counted on invoice date. The official figures and the return come from Serdal.", st["small"])]
    path = os.path.join(a.out, f"Hyperscout_Holding_Maandoverzicht_{m}_{today.isoformat()}.pdf")
    build_pdf(path, "Monthly overview", f"Hyperscout Holding B.V. · {FULL[int(m[5:]) - 1]} {m[:4]}", body, today)
    return [path]

def quarter_pdf(a, data, today):
    fin, fuit, boek = rows(data, "facturen_in", 20), rows(data, "facturen_uit", 18), rows(data, "boekingen", 12)
    q = a.period; y = q[:4]; mon = q_months(q)
    b = vat_boxes(fin, fuit, q)
    def rev(f): return sum(num(r[5]) for r in fuit if f(r[0]))
    def cost(f): return sum(num(r[5]) for r in fin if f(r[0]))
    def bank(f): return -sum(num(r[5]) for r in boek if f(r[0]) and r[6] in ("Bankkosten", "Rente"))
    inq = lambda d: d[:7] in mon; ytd = lambda d: d[:4] == y and d[:7] <= mon[-1]
    res = [["Revenue", eur(rev(inq)), eur(rev(ytd))], ["Costs with invoice", eur(cost(inq)), eur(cost(ytd))],
           ["Bank charges and interest", eur(bank(inq)), eur(bank(ytd))],
           ["Result", eur(rev(inq) - cost(inq) - bank(inq)), eur(rev(ytd) - cost(ytd) - bank(ytd))]]
    miss = missing_rows(boek, upto=mon[-1])
    st = S(); body = []
    body.append(kpis([("BTW to pay (+) / back (−)", eur(b["total"], False)), ("Deadline", deadline(q).strftime("%d %b %Y")),
                      ("Missing invoices", str(len(miss))), ("Owed to Jan (Amex)", eur(num(data.get("amex_owed")), False))]))
    if a.draft:
        body += [Spacer(1, 2 * mm), Paragraph("<b>Draft.</b> The last bank days and late invoices are not in yet. "
                                              "The final version follows on the 15th.", st["body"])]
    for n in data.get("notes") or []:
        body += [Spacer(1, 2 * mm), Paragraph(esc(n), st["body"])]
    body.append(KeepTogether([Paragraph(f"BTW return {q}", st["h2"]), vat_table(b)]))
    body.append(KeepTogether([Paragraph("Result (net of BTW, on invoice date)", st["h2"]),
                              table(["", q, f"Year to date {y}"], res, [94 * mm, 40 * mm, 40 * mm], bold_last=True)]))
    if miss:
        body.append(KeepTogether([Paragraph("Still missing", st["h2"]),
                                  table(["Date", "Source", "Counterparty", "Description", "Amount"], miss,
                                        [22 * mm, 16 * mm, 50 * mm, 60 * mm, 26 * mm], money_from=4)]))
    body += [Spacer(1, 3 * mm), Paragraph("Corporate tax (Vpb) is filed yearly by Serdal; the result above is year to date only, no estimate.", st["small"])]
    tag = "_CONCEPT" if a.draft else ""
    path = os.path.join(a.out, f"Hyperscout_Holding_Kwartaaloverzicht_{q}{tag}_{today.isoformat()}.pdf")
    build_pdf(path, "Quarterly overview" + (" (draft)" if a.draft else ""), f"Hyperscout Holding B.V. · {q}", body, today)
    return [path]

# ------------------------------------------------------------------ Excel for Serdal
def quarter_xlsx(a, data, today):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    fin, fuit, boek = rows(data, "facturen_in", 20), rows(data, "facturen_uit", 18), rows(data, "boekingen", 12)
    q = a.period; mon = q_months(q)
    fin_q = sorted([r for r in fin if quarter_of(r[0]) == q]); fuit_q = sorted([r for r in fuit if quarter_of(r[0]) == q])
    boek_q = sorted([r for r in boek if r[0][:7] in mon]); miss = [r for r in boek if r[8] == "Factuur ontbreekt" and r[0][:7] <= mon[-1]]
    amex_q = [r for r in boek_q if r[1] == "Amex"]
    F = "Arial"; HF = Font(name=F, bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="14213D")
    BOLD = Font(name=F, bold=True); NORM = Font(name=F); BLUE = Font(name=F, color="0000FF")
    EUR = '€ #,##0.00;(€ #,##0.00);"-"'; DT = "yyyy-mm-dd"
    wb = Workbook(); ws0 = wb.active; ws0.title = "Samenvatting"
    def d(s):
        try: return date.fromisoformat(s)
        except Exception: return s
    def sheet(name, header, data_rows, money, dates, widths, formulas=None):
        ws = wb.create_sheet(name)
        ws.append(header)
        for c in ws[1]: c.font, c.fill, c.alignment = HF, HB, Alignment(wrap_text=True, vertical="center")
        for i, r in enumerate(data_rows, start=2):
            for j, v in enumerate(r, start=1):
                cell = ws.cell(row=i, column=j)
                if j in money: cell.value = num(v); cell.number_format = EUR
                elif j in dates: cell.value = d(v) if v else None; cell.number_format = DT
                else: cell.value = v if v != "" else None
                cell.font = NORM
            for j, f in (formulas or {}).items():
                c = ws.cell(row=i, column=j); c.value = f.format(r=i); c.number_format = EUR
        for j, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = "A2"
        return ws
    n_in, n_uit = len(fin_q) + 1, len(fuit_q) + 1
    # Facturen in: I (verlegde btw) and L (bruto) as formulas
    sheet("Facturen in", ["Factuurdatum", "Factuurnr", "Leverancier", "Land", "Op naam Hyperscout", "Netto (€)", "Btw-tarief",
                          "Btw op factuur (€)", "Verlegde btw (€)", "Rubriek", "Aftrekbaar", "Bruto (€)", "Grootboek",
                          "Betaald via", "Betaald op", "Bestand", "Valuta", "Bedrag in valuta", "Notitie"],
          [r[:8] + [""] + r[9:11] + [""] + r[12:16] + r[17:20] for r in fin_q], {6, 8}, {1, 15},
          [12, 14, 26, 6, 10, 13, 14, 13, 13, 8, 10, 13, 24, 12, 12, 40, 7, 12, 30],
          {9: "=IF(LEFT(G{r},7)=\"verlegd\",ROUND(F{r}*'BTW-aangifte'!$C$3,2),0)", 12: "=F{r}+H{r}"})
    sheet("Facturen uit", ["Factuurdatum", "Factuurnr", "Klant", "Land", "Btw-nummer klant", "Netto (€)", "Btw-tarief", "Btw (€)",
                           "Bruto (€)", "Rubriek", "Vervaldatum", "Betaald op", "Bestand", "Valuta", "Bedrag in valuta", "Notitie"],
          [r[:8] + [""] + r[9:12] + [r[13]] + r[15:18] for r in fuit_q], {6, 8}, {1, 11, 12},
          [12, 14, 26, 6, 18, 13, 10, 13, 13, 9, 12, 12, 40, 7, 12, 30], {9: "=F{r}+H{r}"})
    ws = wb.create_sheet("BTW-aangifte", 1)
    ws["A1"] = f"Btw-aangifte {q}, Hyperscout Holding B.V."; ws["A1"].font = Font(name=F, bold=True, size=13)
    ws["A3"] = "Btw hoog tarief"; ws["C3"] = VAT_HIGH; ws["C3"].number_format = "0%"; ws["C3"].font = BLUE
    ws["D3"] = "Bron: belastingdienst.nl"
    ws.append([]); ws.append(["Rubriek", "Omschrijving", "Omzet / grondslag (€)", "Btw (€)"])
    for c in ws[5]: c.font, c.fill = HF, HB
    U, I = "'Facturen uit'!", "'Facturen in'!"
    lines = [("1a", "Leveringen/diensten belast met hoog tarief", f"=SUMIFS({U}F:F,{U}J:J,\"1a\")", f"=SUMIFS({U}H:H,{U}J:J,\"1a\")"),
             ("1b", "Leveringen/diensten belast met laag tarief", f"=SUMIFS({U}F:F,{U}J:J,\"1b\")", f"=SUMIFS({U}H:H,{U}J:J,\"1b\")"),
             ("3a", "Leveringen naar landen buiten de EU", f"=SUMIFS({U}F:F,{U}J:J,\"3a\")", None),
             ("3b", "Leveringen naar of diensten in landen binnen de EU", f"=SUMIFS({U}F:F,{U}J:J,\"3b\")", None),
             ("4a", "Leveringen/diensten uit landen buiten de EU", f"=SUMIFS({I}F:F,{I}J:J,\"4a\")", f"=SUMIFS({I}I:I,{I}J:J,\"4a\")"),
             ("4b", "Leveringen/diensten uit landen binnen de EU", f"=SUMIFS({I}F:F,{I}J:J,\"4b\")", f"=SUMIFS({I}I:I,{I}J:J,\"4b\")"),
             ("5a", "Verschuldigde btw", None, "=D6+D7+D10+D11"),
             ("5b", "Voorbelasting", None, f"=SUMIFS({I}H:H,{I}K:K,\"Ja\")+SUMIFS({I}I:I,{I}K:K,\"Ja\")"),
             ("", "Te betalen (+) of terug te vragen (−)", None, "=D12-D13")]
    for k, t, base, btw in lines:
        ws.append([k, t, base, btw])
        for c in ws[ws.max_row][2:]: c.number_format = EUR
        for c in ws[ws.max_row]: c.font = NORM
    for c in ws[ws.max_row]: c.font = BOLD
    ws.append([]); ws.append(["", "Uiterste aangifte- en betaaldatum", deadline(q)]); ws.cell(ws.max_row, 3).number_format = DT
    ws.append(["", "Niet in de aangifte: diensten aan klanten buiten de EU", f"=SUMIFS({U}F:F,{U}J:J,\"buiten EU\")"])
    ws.cell(ws.max_row, 3).number_format = EUR
    ws.append(["", "Btw op inkoop die niet aftrekbaar is", f"=SUMIFS({I}H:H,{I}K:K,\"Nee\")"]); ws.cell(ws.max_row, 3).number_format = EUR
    for col, w in zip("ABCD", (9, 52, 22, 16)): ws.column_dimensions[col].width = w
    icp = wb.create_sheet("ICP-opgave", 2)
    icp.append(["Klant", "Land", "Btw-nummer", "Bedrag diensten (€)"])
    for c in icp[1]: c.font, c.fill = HF, HB
    per = {}
    for r in fuit_q:
        if r[9] == "3b": per.setdefault((r[2], r[3], r[4]), 0.0); per[(r[2], r[3], r[4])] += num(r[5])
    for (k, l, v), s in sorted(per.items()):
        icp.append([k, l, v, s]); icp.cell(icp.max_row, 4).number_format = EUR
    icp.append(["Totaal", "", "", f"=SUM(D2:D{max(icp.max_row, 2)})"]); icp.cell(icp.max_row, 4).number_format = EUR
    for c in icp[icp.max_row]: c.font = BOLD
    icp.append([]); icp.append(["Moet aansluiten op rubriek 3b in de tab BTW-aangifte.", "", "", "=D" + str(icp.max_row - 1) + "-'BTW-aangifte'!C9"])
    icp.cell(icp.max_row, 4).number_format = EUR
    for col, w in zip("ABCD", (34, 8, 20, 20)): icp.column_dimensions[col].width = w
    bk = ["Datum", "Bron", "Referentie", "Tegenpartij", "Omschrijving", "Bedrag (€)", "Grootboek", "Factuurnr", "Status", "Notitie"]
    sheet("Boekingen", bk, [r[:9] + [r[11]] for r in boek_q], {6}, {1}, [12, 8, 16, 26, 30, 13, 24, 14, 16, 30])
    sheet("Ontbrekend", ["Datum", "Bron", "Tegenpartij", "Omschrijving", "Bedrag (€)", "Notitie"],
          [[r[0], r[1], r[3], r[4], r[5], r[11]] for r in sorted(miss)], {5}, {1}, [12, 8, 26, 30, 13, 30])
    sheet("Amex declaraties", ["Datum", "Tegenpartij", "Omschrijving", "Bedrag (€)", "Factuurnr", "Status"],
          [[r[0], r[3], r[4], r[5], r[7], r[8]] for r in amex_q], {4}, {1}, [12, 26, 30, 13, 14, 16])
    # Samenvatting (formulas over the tabs)
    s = ws0
    s["A1"] = f"Hyperscout Holding B.V., boekhouding {q}" + (" (CONCEPT)" if a.draft else ""); s["A1"].font = Font(name=F, bold=True, size=14)
    s["A2"] = f"Opgemaakt {today.isoformat()} voor Serdal. Bedragen op factuurdatum, exclusief btw tenzij anders vermeld."
    s["A2"].font = Font(name=F, italic=True, color="5F6368")
    tbl = [("Omzet kwartaal (€)", "=SUM('Facturen uit'!F:F)"),
           ("Kosten met factuur kwartaal (€)", "=SUM('Facturen in'!F:F)"),
           ("Bankkosten en rente kwartaal (€)", "=-SUMIFS(Boekingen!F:F,Boekingen!G:G,\"Bankkosten\")-SUMIFS(Boekingen!F:F,Boekingen!G:G,\"Rente\")"),
           ("Resultaat kwartaal (€)", "=B4-B5-B6"), (None, None),
           ("Btw te betalen (+) of terug (−) (€)", "='BTW-aangifte'!D14"),
           ("Uiterste aangifte- en betaaldatum", deadline(q)), (None, None),
           ("Boekingen zonder factuur", "=MAX(COUNTA(Ontbrekend!A:A)-1,0)"),
           ("Bedrag zonder factuur (€)", "=-SUM(Ontbrekend!E:E)"),
           ("Amex-voorschotten Jan dit kwartaal (€)", "=-SUM('Amex declaraties'!D:D)"),
           ("Nog te betalen aan Jan, Amex totaal (€)", num(data.get("amex_owed")))]
    for i, (l, f) in enumerate(tbl, start=4):
        if l is None: continue
        s.cell(i, 1, l).font = NORM; c = s.cell(i, 2, f); c.font = NORM
        c.number_format = DT if isinstance(f, date) else ("0" if l.startswith("Boekingen zonder") else EUR)
    s["A7"].font = BOLD; s["B7"].font = BOLD; s["A9"].font = BOLD; s["B9"].font = BOLD
    s.cell(15, 3, "Bron: ledger 'Amex declaraties', stand op uitgiftedatum")
    s.column_dimensions["A"].width = 44; s.column_dimensions["B"].width = 18; s.column_dimensions["C"].width = 44
    for wsx in wb.worksheets:
        for row in wsx.iter_rows():
            for c in row:
                if c.font is None or c.font.name != F: c.font = Font(name=F, bold=c.font.bold if c.font else False,
                                                                         color=c.font.color if c.font else None)
    tag = "_CONCEPT" if a.draft else ""
    path = os.path.join(a.out, f"Hyperscout_Holding_BTW_{q}{tag}_{today.isoformat()}.xlsx")
    wb.save(path)
    return [path]

if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--type", required=True, choices=["month", "quarter"])
    p.add_argument("--period", required=True, help="2026-09 for month, 2026-Q3 for quarter")
    p.add_argument("--data", required=True); p.add_argument("--out", default=".")
    p.add_argument("--draft", action="store_true"); p.add_argument("--date", default=None)
    a = p.parse_args(); data = json.load(open(a.data))
    today = date.fromisoformat(a.date) if a.date else date.today()
    files = month_pdf(a, data, today) if a.type == "month" else quarter_pdf(a, data, today) + quarter_xlsx(a, data, today)
    print("\n".join(files))
